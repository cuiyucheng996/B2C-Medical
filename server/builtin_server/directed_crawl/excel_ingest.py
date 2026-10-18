"""影刀导出 Excel：BSR / 单品深度采集（见 doc/External/*.md）。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _sheet_rows(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        item: dict[str, Any] = {}
        for i, key in enumerate(header):
            if not key:
                continue
            val = row[i] if i < len(row) else None
            item[key] = val
        if item:
            out.append(item)
    return out


def _read_sheet(wb, name: str):
    if name in wb.sheetnames:
        return wb[name]
    for sn in wb.sheetnames:
        if sn.strip() == name.strip():
            return wb[sn]
    return None


def ingest_bsr_excel(path: Path, *, main_sheet: str = "BSR100", ids_sheet: str = "RANK&ASIN") -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheetnames = list(wb.sheetnames)
    main_ws = _read_sheet(wb, main_sheet) or wb.active
    main_rows = _sheet_rows(main_ws) if main_ws else []
    ids_ws = _read_sheet(wb, ids_sheet)
    id_rows = _sheet_rows(ids_ws) if ids_ws else []
    wb.close()
    pairs: list[dict[str, Any]] = []
    for row in id_rows:
        asin = row.get("ASIN") or row.get("asin")
        rank = row.get("RANK") or row.get("Rank") or row.get("rank")
        if asin is not None or rank is not None:
            pairs.append({"asin": asin, "rank": rank})
    return {
        "kind": "amazon_bsr",
        "main_count": len(main_rows),
        "id_count": len(pairs),
        "main_preview": main_rows[:5],
        "rank_asin": pairs[:200],
        "sheets_found": sheetnames,
    }


def ingest_product_excel(
    path: Path,
    *,
    goods_sheet: str = "商品信息",
    remarks_sheet: str = "评论数据",
) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    goods_ws = _read_sheet(wb, goods_sheet)
    remarks_ws = _read_sheet(wb, remarks_sheet)
    goods = _sheet_rows(goods_ws) if goods_ws else []
    remarks = _sheet_rows(remarks_ws) if remarks_ws else []
    sheetnames = list(wb.sheetnames)
    wb.close()
    return {
        "kind": "amazon_product",
        "goods_count": len(goods),
        "remarks_count": len(remarks),
        "goods_preview": goods[:3],
        "remarks_preview": remarks[:5],
        "sheets_found": sheetnames,
    }


def ingest_excel_auto(path: Path, job_type: str = "") -> dict[str, Any]:
    kind = (job_type or "").strip().lower()
    if kind in {"bsr", "amazon_bsr", "bst"}:
        return ingest_bsr_excel(path)
    if kind in {"product", "amazon_product", "detail"}:
        return ingest_product_excel(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    if "RANK&ASIN" in names or "BSR100" in names:
        return ingest_bsr_excel(path)
    if "商品信息" in names or "评论数据" in names:
        return ingest_product_excel(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = _sheet_rows(ws) if ws else []
    wb.close()
    return {"kind": "generic", "row_count": len(rows), "preview": rows[:10], "sheets_found": names}
